"""Service layer for superadmin module."""

import uuid
from datetime import UTC, datetime, timedelta, timezone

import jwt as pyjwt
from sqlalchemy import func, select, text

from src.common.enums import AuditAction, PlanTier, SubscriptionStatus, UserRole
from src.config import settings
from src.infrastructure.uow import UnitOfWork
from src.modules.ai.models import AIUsageLog
from src.modules.audit.repository import AuditRepository
from src.modules.auth.models import User
from src.modules.files.models import File
from src.modules.org.models import Membership, Organization, Subscription
from src.modules.tables.models import Table
from src.modules.tables.records import Record
from src.modules.superadmin.repository import SuperadminRepository
from src.modules.tables.repository import TableRepository
from src.modules.tables.records import RecordRepository


def authenticate_superadmin(email: str, password: str) -> str:
    if not settings.SUPERADMIN_EMAIL or not settings.SUPERADMIN_PASSWORD:
        raise RuntimeError("SUPERADMIN_NOT_CONFIGURED")
    if email != settings.SUPERADMIN_EMAIL or password != settings.SUPERADMIN_PASSWORD:
        raise ValueError("INVALID_CREDENTIALS")
    now = datetime.now(UTC)
    payload = {
        "sub": "superadmin",
        "role": UserRole.SUPERADMIN.value,
        "type": "access",
        "iat": now,
        "exp": now + timedelta(hours=12),
    }
    # Must match auth/security.py decode_access_token algorithm.
    return pyjwt.encode(payload, settings.SECRET_KEY, algorithm=settings.JWT_ALGORITHM)


def _safe_int(v) -> int:
    try:
        return int(v or 0)
    except Exception:
        return 0


async def dashboard_data() -> dict:
    async with UnitOfWork() as uow:
        orgs_count = (await uow.session.execute(select(func.count()).select_from(Organization))).scalar() or 0
        users_count = (await uow.session.execute(select(func.count()).select_from(User))).scalar() or 0
        tables_count = (await uow.session.execute(select(func.count()).select_from(Table))).scalar() or 0
        records_count = (await uow.session.execute(select(func.count()).select_from(Record))).scalar() or 0
        files_row = (
            await uow.session.execute(select(func.count(), func.coalesce(func.sum(File.size), 0)).select_from(File))
        ).one()
        ai_row = (
            await uow.session.execute(select(func.count(), func.coalesce(func.sum(AIUsageLog.total_tokens), 0)).select_from(AIUsageLog))
        ).one()

        thirty_days_ago = datetime.now(timezone.utc) - timedelta(days=30)
        reg_stmt = (
            select(func.date_trunc("day", User.created_at).label("day"), func.count().label("cnt"))
            .where(User.created_at >= thirty_days_ago)
            .group_by(text("day"))
            .order_by(text("day"))
        )
        reg_rows = (await uow.session.execute(reg_stmt)).all()

        plan_stmt = select(Organization.plan, func.count().label("cnt")).group_by(Organization.plan)
        plan_rows = (await uow.session.execute(plan_stmt)).all()

    return {
        "totals": {
            "orgs": _safe_int(orgs_count),
            "users": _safe_int(users_count),
            "tables": _safe_int(tables_count),
            "records": _safe_int(records_count),
            "files": _safe_int(files_row[0]),
            "storage_bytes": _safe_int(files_row[1]),
            "ai_requests": _safe_int(ai_row[0]),
            "ai_tokens": _safe_int(ai_row[1]),
        },
        "registrations_timeline": [{"date": str(r.day)[:10], "count": r.cnt} for r in reg_rows],
        "orgs_by_plan": [{"plan": str(r.plan.value if hasattr(r.plan, "value") else r.plan), "count": r.cnt} for r in plan_rows],
    }


async def list_orgs_page(*, q: str | None, plan: str | None, sub_status: str | None, limit: int, offset: int) -> dict:
    async with UnitOfWork() as uow:
        repo = SuperadminRepository(uow.session)
        items, total = await repo.list_orgs_page(q=q, plan=plan, sub_status=sub_status, limit=limit, offset=offset)
    return {"items": items, "total": total, "limit": int(limit), "offset": int(offset)}


async def org_detail(org_id: str) -> dict:
    async with UnitOfWork() as uow:
        repo = SuperadminRepository(uow.session)
        data = await repo.get_org_detail(org_id=uuid.UUID(org_id))
        if not data:
            raise LookupError("NOT_FOUND")
    return data


async def org_members_page(org_id: str, *, limit: int, offset: int) -> dict:
    async with UnitOfWork() as uow:
        repo = SuperadminRepository(uow.session)
        items, total = await repo.list_org_members(org_id=uuid.UUID(org_id), limit=limit, offset=offset)
    return {"items": items, "total": total, "limit": int(limit), "offset": int(offset)}


async def list_users_page(
    *,
    q: str | None,
    org_id: str | None,
    is_active: bool | None,
    limit: int,
    offset: int,
) -> dict:
    async with UnitOfWork() as uow:
        repo = SuperadminRepository(uow.session)
        items, total = await repo.list_users_page(
            q=q,
            org_id=uuid.UUID(org_id) if org_id else None,
            is_active=is_active,
            limit=limit,
            offset=offset,
        )
    return {"items": items, "total": total, "limit": int(limit), "offset": int(offset)}


async def list_tables(limit: int, offset: int) -> list[dict]:
    async with UnitOfWork() as uow:
        stmt = select(Table).order_by(Table.created_at.desc()).limit(limit).offset(offset)
        tables = list((await uow.session.execute(stmt)).scalars().all())
        result: list[dict] = []
        for table in tables:
            rec_cnt = (
                await uow.session.execute(select(func.count()).select_from(Record).where(Record.table_id == table.id))
            ).scalar() or 0
            result.append(
                {
                    "id": str(table.id),
                    "name": table.name,
                    "org_id": str(table.org_id),
                    "columns": len(table.columns) if table.columns else 0,
                    "records": rec_cnt,
                    "created_at": table.created_at.isoformat() if table.created_at else None,
                }
            )
    return result


async def ai_usage_by_org() -> list[dict]:
    async with UnitOfWork() as uow:
        stmt = (
            select(
                AIUsageLog.org_id,
                func.count().label("requests"),
                func.coalesce(func.sum(AIUsageLog.total_tokens), 0).label("tokens"),
            )
            .group_by(AIUsageLog.org_id)
            .order_by(func.sum(AIUsageLog.total_tokens).desc())
        )
        rows = (await uow.session.execute(stmt)).all()
        org_ids = [r.org_id for r in rows]
        orgs_map: dict = {}
        if org_ids:
            orgs = list((await uow.session.execute(select(Organization).where(Organization.id.in_(org_ids)))).scalars().all())
            orgs_map = {o.id: o.name for o in orgs}
    return [
        {
            "org_id": str(r.org_id),
            "org_name": orgs_map.get(r.org_id, "—"),
            "requests": _safe_int(r.requests),
            "tokens": _safe_int(r.tokens),
        }
        for r in rows
    ]


async def set_plan(org_id: str, plan_name: str) -> dict:
    try:
        plan_tier = PlanTier(plan_name)
    except ValueError as exc:
        raise ValueError("INVALID_PLAN") from exc

    async with UnitOfWork() as uow:
        org_uuid = uuid.UUID(org_id)
        org = (await uow.session.execute(select(Organization).where(Organization.id == org_uuid))).scalar_one_or_none()
        if not org:
            raise LookupError("NOT_FOUND")
        old_plan = org.plan.value if hasattr(org.plan, "value") else str(org.plan)

        # Update subscription (source of truth for many checks) and keep org.plan in sync (legacy).
        sub = (await uow.session.execute(select(Subscription).where(Subscription.org_id == org_uuid))).scalar_one_or_none()
        if sub:
            sub.plan = plan_tier
            sub.status = SubscriptionStatus.ACTIVE
        else:
            sub = Subscription(org_id=org_uuid, plan=plan_tier, status=SubscriptionStatus.ACTIVE)
            uow.session.add(sub)

        org.plan = plan_tier

        audit_repo = AuditRepository(uow.session)
        await audit_repo.log(
            org_id=org_uuid,
            actor_id=None,
            action=AuditAction.UPDATE,
            entity_type="org_plan",
            entity_id=str(org_uuid),
            meta={"superadmin": True, "old_plan": old_plan, "new_plan": plan_tier.value},
        )
        await uow.commit()
    return {"org_id": org_id, "plan": plan_name}


def ai_config() -> dict:
    base_url = (settings.AI_BASE_URL or "").rstrip("/")
    provider = "timeweb-agent-openai-compatible" if "agent.timeweb.cloud" in base_url else "openai-compatible"
    key = settings.OPENAI_BEARER_TOKEN or settings.OPENAI_API_KEY or ""
    return {
        "provider": provider,
        "base_url": base_url,
        "official_provider_docs_url": "https://agent.timeweb.cloud/docs",
        "model": settings.OPENAI_MODEL,
        "key_configured": bool(key),
        "key_prefix": f"{key[:4]}***" if key else "",
    }


async def list_org_options(limit: int = 200) -> list[dict]:
    """Minimal org list for a superadmin selector (no heavy per-org counts)."""
    async with UnitOfWork() as uow:
        stmt = select(Organization).order_by(Organization.created_at.desc()).limit(limit)
        orgs = list((await uow.session.execute(stmt)).scalars().all())
        return [
            {
                "id": str(o.id),
                "name": o.name,
                "slug": o.slug,
                "plan": o.plan.value if hasattr(o.plan, "value") else str(o.plan),
                "created_at": o.created_at.isoformat() if o.created_at else None,
            }
            for o in orgs
        ]


async def overview_data(org_limit: int = 200) -> dict:
    """Dashboard + org selector options for the superadmin UI."""
    dash = await dashboard_data()
    orgs = await list_org_options(limit=org_limit)
    return {
        "dashboard": dash,
        "orgs": orgs,
        "generated_at": datetime.now(UTC).isoformat(),
    }


async def audit_logs_page(*, org_id: str | None, limit: int, offset: int) -> dict:
    async with UnitOfWork() as uow:
        repo = SuperadminRepository(uow.session)
        items, total = await repo.list_audit_logs_page(
            org_id=uuid.UUID(org_id) if org_id else None,
            limit=limit,
            offset=offset,
        )
    return {"items": items, "total": total, "limit": int(limit), "offset": int(offset)}


async def org_tables_page(
    org_id: str,
    *,
    q: str | None,
    include_archived: bool,
    limit: int,
    offset: int,
) -> dict:
    async with UnitOfWork() as uow:
        repo = SuperadminRepository(uow.session)
        items, total = await repo.list_tables_by_org_page(
            org_id=uuid.UUID(org_id),
            q=q,
            include_archived=include_archived,
            limit=limit,
            offset=offset,
        )
    return {"items": items, "total": total, "limit": int(limit), "offset": int(offset)}


async def org_table_detail(org_id: str, table_id: str) -> dict:
    async with UnitOfWork() as uow:
        repo = SuperadminRepository(uow.session)
        table = await repo.get_table_in_org(org_id=uuid.UUID(org_id), table_id=uuid.UUID(table_id))
        if not table:
            raise LookupError("NOT_FOUND")
        # Serialize minimal table with columns for read-only UI.
        columns = [
            {
                "id": str(c.id),
                "name": c.name,
                "field_type": c.field_type,
                "position": int(c.position),
                "is_required": bool(c.is_required),
                "is_primary": bool(c.is_primary),
                "config": c.config,
                "default_value": c.default_value,
            }
            for c in sorted(table.columns or [], key=lambda x: x.position)
        ]
        return {
            "id": str(table.id),
            "org_id": str(table.org_id),
            "folder_id": str(table.folder_id) if table.folder_id else None,
            "name": table.name,
            "description": table.description,
            "icon": table.icon,
            "color": table.color,
            "is_archived": bool(table.is_archived),
            "created_at": table.created_at.isoformat() if table.created_at else None,
            "columns": columns,
        }


async def org_table_records_page(
    org_id: str,
    table_id: str,
    *,
    q: str | None,
    sort_col_id: str | None,
    sort_dir: str,
    limit: int,
    offset: int,
) -> dict:
    if sort_dir not in {"asc", "desc"}:
        sort_dir = "asc"
    async with UnitOfWork() as uow:
        repo = SuperadminRepository(uow.session)
        # Ensure table belongs to org.
        table = await repo.get_table_in_org(org_id=uuid.UUID(org_id), table_id=uuid.UUID(table_id))
        if not table:
            raise LookupError("NOT_FOUND")
        rows, total = await repo.list_table_records_page(
            org_id=uuid.UUID(org_id),
            table_id=uuid.UUID(table_id),
            q=q,
            sort_col_id=sort_col_id,
            sort_dir=sort_dir,
            limit=limit,
            offset=offset,
        )
        items = []
        for r in rows:
            items.append(
                {
                    "id": str(r.id),
                    "table_id": str(r.table_id),
                    "data": r.data,
                    "created_by": str(r.created_by) if r.created_by else None,
                    "created_at": r.created_at.isoformat() if r.created_at else None,
                    "updated_at": r.updated_at.isoformat() if r.updated_at else None,
                    "position": int(r.position),
                }
            )
    return {"items": items, "total": total, "limit": int(limit), "offset": int(offset)}


async def export_table_csv(*, org_id: str, table_id: str) -> tuple[bytes, str]:
    """Экспорт таблицы в CSV (UTF-8 with BOM) для supеradmin."""
    import csv
    import io

    async with UnitOfWork() as uow:
        t_repo = TableRepository(uow.session)
        table_uuid = uuid.UUID(table_id)
        org_uuid = uuid.UUID(org_id)

        table = await t_repo.get_by_id(table_uuid, with_columns=True)
        if not table or table.org_id != org_uuid:
            raise LookupError("NOT_FOUND")

        r_repo = RecordRepository(uow.session)
        records = await r_repo.list_by_table(table_uuid, limit=5000, offset=0)

        columns = sorted(table.columns, key=lambda c: c.position)
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([c.name for c in columns])
        for rec in records:
            writer.writerow([str(rec.data.get(str(c.id), "")) for c in columns])

        payload = output.getvalue().encode("utf-8-sig")

        audit_repo = AuditRepository(uow.session)
        await audit_repo.log(
            org_id=org_uuid,
            actor_id=None,
            action=AuditAction.EXPORT,
            entity_type="table_export",
            entity_id=str(table_uuid),
            meta={"superadmin": True, "format": "csv"},
        )
        await uow.commit()

        filename = f"{table.name}.csv"
        return payload, filename


async def export_table_xlsx(*, org_id: str, table_id: str) -> tuple[bytes, str]:
    """Экспорт таблицы в XLSX для supеradmin."""
    from io import BytesIO

    from openpyxl import Workbook

    async with UnitOfWork() as uow:
        t_repo = TableRepository(uow.session)
        table_uuid = uuid.UUID(table_id)
        org_uuid = uuid.UUID(org_id)

        table = await t_repo.get_by_id(table_uuid, with_columns=True)
        if not table or table.org_id != org_uuid:
            raise LookupError("NOT_FOUND")

        r_repo = RecordRepository(uow.session)
        records = await r_repo.list_by_table(table_uuid, limit=5000, offset=0)

        columns = sorted(table.columns, key=lambda c: c.position)
        wb = Workbook()
        ws = wb.active
        ws.title = "Table"
        ws.append([c.name for c in columns])
        for rec in records:
            ws.append([str(rec.data.get(str(c.id), "")) for c in columns])

        output = BytesIO()
        wb.save(output)
        payload = output.getvalue()

        audit_repo = AuditRepository(uow.session)
        await audit_repo.log(
            org_id=org_uuid,
            actor_id=None,
            action=AuditAction.EXPORT,
            entity_type="table_export",
            entity_id=str(table_uuid),
            meta={"superadmin": True, "format": "xlsx"},
        )
        await uow.commit()

        filename = f"{table.name}.xlsx"
        return payload, filename
