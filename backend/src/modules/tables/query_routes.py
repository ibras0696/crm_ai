"""Record filter/search and table export/import endpoints."""

from __future__ import annotations

import csv
import io
import uuid
from io import BytesIO

from fastapi import APIRouter, Depends, File, Query, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook

from src.common.enums import UserRole
from src.common.http_headers import content_disposition_attachment
from src.common.schemas import ApiResponse
from src.infrastructure.metrics_custom import EXPORTS_TOTAL, IMPORTS_TOTAL
from src.infrastructure.uow import UnitOfWork
from src.modules.access.dependencies import require_access
from src.modules.auth.dependencies import CurrentUser, require_roles
from src.modules.tables.repository import TableRepository
from src.modules.tables.records import Record, RecordRepository
from src.modules.tables.schemas import FilterRequest, RecordOut


router = APIRouter(prefix="/tables/{table_id}", tags=["records"])


@router.post("/filter", response_model=ApiResponse[dict])
async def filter_records(
    table_id: uuid.UUID,
    body: FilterRequest,
    limit: int = Query(default=100, le=500),
    offset: int = Query(default=0, ge=0),
    current_user: CurrentUser = Depends(
        require_roles(UserRole.OWNER, UserRole.ADMIN, UserRole.MANAGER, UserRole.EMPLOYEE, UserRole.READONLY)
    ),
    _: None = Depends(require_access(resource_type="table", permission="can_read", resource_id_param="table_id")),
):
    """Filter records. Current implementation is in-memory, intended for small datasets."""
    async with UnitOfWork() as uow:
        t_repo = TableRepository(uow.session)
        table = await t_repo.get_by_id(table_id, with_columns=False)
        if not table or table.org_id != current_user.org_id:
            return ApiResponse(ok=False, data=None, error={"code": "NOT_FOUND", "message": "Таблица не найдена"})

        r_repo = RecordRepository(uow.session)
        all_records = await r_repo.list_by_table(table_id, limit=500, offset=0)

        filtered = all_records
        if body.filters:
            for col_id, cond in body.filters.items():
                op = cond.get("op", "eq") if isinstance(cond, dict) else "eq"
                val = cond.get("value", "") if isinstance(cond, dict) else cond
                val_str = str(val).lower()
                new_filtered = []
                for rec in filtered:
                    cell = str(rec.data.get(col_id, "")).lower()
                    if op == "eq" and cell == val_str:
                        new_filtered.append(rec)
                    elif op == "contains" and val_str in cell:
                        new_filtered.append(rec)
                    elif op == "gt" and cell > val_str:
                        new_filtered.append(rec)
                    elif op == "lt" and cell < val_str:
                        new_filtered.append(rec)
                    elif op == "neq" and cell != val_str:
                        new_filtered.append(rec)
                filtered = new_filtered

        if body.sorts:
            for s in reversed(body.sorts):
                col_id = s.get("col_id", "")
                desc = s.get("dir", "asc") == "desc"
                filtered.sort(key=lambda r: str(r.data.get(col_id, "")), reverse=desc)

        total = len(filtered)
        page = filtered[offset : offset + limit]
        items = [RecordOut.model_validate(r) for r in page]
    return ApiResponse(data={"records": [i.model_dump() for i in items], "total": total})


@router.get("/export/csv")
async def export_csv(
    table_id: uuid.UUID,
    current_user: CurrentUser = Depends(require_roles(UserRole.OWNER, UserRole.ADMIN, UserRole.MANAGER, UserRole.EMPLOYEE)),
    _: None = Depends(require_access(resource_type="table", permission="can_read", resource_id_param="table_id")),
):
    async with UnitOfWork() as uow:
        t_repo = TableRepository(uow.session)
        table = await t_repo.get_by_id(table_id, with_columns=True)
        if not table or table.org_id != current_user.org_id:
            return ApiResponse(ok=False, data=None, error={"code": "NOT_FOUND", "message": "Таблица не найдена"})

        r_repo = RecordRepository(uow.session)
        records = await r_repo.list_by_table(table_id, limit=5000, offset=0)

        columns = sorted(table.columns, key=lambda c: c.position)
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([c.name for c in columns])
        for rec in records:
            writer.writerow([str(rec.data.get(str(c.id), "")) for c in columns])

        payload = output.getvalue().encode("utf-8-sig")
        try:
            EXPORTS_TOTAL.labels(format="csv").inc()
        except Exception:
            pass

        return StreamingResponse(
            iter([payload]),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": content_disposition_attachment(f"{table.name}.csv")},
        )


@router.get("/export/xlsx")
async def export_xlsx(
    table_id: uuid.UUID,
    current_user: CurrentUser = Depends(require_roles(UserRole.OWNER, UserRole.ADMIN, UserRole.MANAGER, UserRole.EMPLOYEE)),
    _: None = Depends(require_access(resource_type="table", permission="can_read", resource_id_param="table_id")),
):
    async with UnitOfWork() as uow:
        t_repo = TableRepository(uow.session)
        table = await t_repo.get_by_id(table_id, with_columns=True)
        if not table or table.org_id != current_user.org_id:
            return ApiResponse(ok=False, data=None, error={"code": "NOT_FOUND", "message": "Таблица не найдена"})

        r_repo = RecordRepository(uow.session)
        records = await r_repo.list_by_table(table_id, limit=5000, offset=0)

        columns = sorted(table.columns, key=lambda c: c.position)
        wb = Workbook()
        ws = wb.active
        ws.title = "Table"
        ws.append([c.name for c in columns])
        for rec in records:
            ws.append([str(rec.data.get(str(c.id), "")) for c in columns])

        output = BytesIO()
        wb.save(output)
        payload = output.getvalue()

        try:
            EXPORTS_TOTAL.labels(format="xlsx").inc()
        except Exception:
            pass

        return StreamingResponse(
            iter([payload]),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": content_disposition_attachment(f"{table.name}.xlsx")},
        )


@router.post("/import/csv", response_model=ApiResponse[dict])
async def import_csv(
    table_id: uuid.UUID,
    file: UploadFile = File(...),
    current_user: CurrentUser = Depends(require_roles(UserRole.OWNER, UserRole.ADMIN, UserRole.MANAGER)),
    _: None = Depends(require_access(resource_type="table", permission="can_write", resource_id_param="table_id")),
):
    """
    Import CSV into an existing table (basic version).

    Behavior:
    - matches CSV header columns to existing table columns by name (case-insensitive).
    - creates records for matched columns only.
    """
    async with UnitOfWork() as uow:
        t_repo = TableRepository(uow.session)
        table = await t_repo.get_by_id(table_id, with_columns=True)
        if not table or table.org_id != current_user.org_id:
            return ApiResponse(ok=False, data=None, error={"code": "NOT_FOUND", "message": "Таблица не найдена"})

        raw = await file.read()
        text = raw.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        try:
            header = next(reader)
        except StopIteration:
            return ApiResponse(ok=False, data=None, error={"code": "BAD_REQUEST", "message": "Пустой CSV"})

        by_name = {str(c.name).strip().lower(): c for c in table.columns}
        col_map: list[tuple[int, str]] = []
        for idx, name in enumerate(header):
            col = by_name.get(str(name).strip().lower())
            if col:
                col_map.append((idx, str(col.id)))

        if not col_map:
            return ApiResponse(
                ok=False,
                data=None,
                error={"code": "NO_MATCHING_COLUMNS", "message": "В CSV не найдено совпадений с колонками таблицы."},
            )

        r_repo = RecordRepository(uow.session)
        created = 0
        max_pos = await r_repo.get_max_position(table_id)
        to_create: list[Record] = []
        for row in reader:
            data: dict[str, str] = {}
            for idx, col_id in col_map:
                if idx < len(row):
                    v = row[idx]
                    if v is not None and str(v).strip() != "":
                        data[col_id] = str(v)
            if not data:
                continue
            to_create.append(
                Record(
                    table_id=table_id,
                    org_id=current_user.org_id,
                    created_by=current_user.user_id,
                    data=data,
                    position=max_pos + created + 1,
                )
            )
            created += 1
            if created >= 2000:
                break

        if to_create:
            await r_repo.bulk_create(to_create)

        await uow.commit()

        try:
            IMPORTS_TOTAL.labels(format="csv").inc()
        except Exception:
            pass

    return ApiResponse(data={"records_created": created})
