"""Сервис фильтрации/экспорта/импорта записей таблиц.

Здесь бизнес-логика. Роуты должны быть тонкими: валидация входа + вызов сервиса.
"""

from __future__ import annotations

import csv
import io
import json
import time
from contextlib import suppress
from datetime import date, datetime
from io import BytesIO
from typing import TYPE_CHECKING

from openpyxl import Workbook
from sqlalchemy import Date, DateTime, Float, and_, cast, func, or_, select

from src.common.enums import AuditAction
from src.config import settings
from src.infrastructure.metrics_custom import EXPORTS_TOTAL, IMPORTS_TOTAL
from src.modules.audit.models import AuditLog
from src.modules.tables.records import Record, RecordRepository
from src.modules.tables.repository import TablePlanLimitsRepository, TableRepository
from src.modules.tables.services import TableFormulaEngine, TableRelationEngine

if TYPE_CHECKING:
    import uuid


class TableQueryService:
    """Операции над записями (read-heavy): фильтры/экспорт/импорт."""

    def __init__(self, session):
        self.session = session
        self.t_repo = TableRepository(session)
        self.r_repo = RecordRepository(session)
        self.plan_limits_repo = TablePlanLimitsRepository(session)
        self.relation_engine = TableRelationEngine(session)
        self.formula_engine = TableFormulaEngine()

    async def filter_records(
        self,
        *,
        table_id: uuid.UUID,
        org_id: uuid.UUID,
        search: str | None,
        filters: dict | list | None,
        sorts: list | None,
        limit: int,
        offset: int,
    ) -> dict:
        """Фильтрация/сортировка записей с пагинацией."""
        table = await self.t_repo.get_by_id(table_id, with_columns=True)
        if not table or table.org_id != org_id:
            raise LookupError("NOT_FOUND")

        columns = {str(col.id): col for col in table.columns}
        stmt = select(Record).where(Record.table_id == table_id)

        parsed_filters = self._normalize_filters(filters)
        for item in parsed_filters:
            condition = self._build_filter_condition(item=item, columns=columns)
            if condition is not None:
                stmt = stmt.where(condition)

        if parsed_filters:
            pass

        # Optional global search across all columns.
        search_value = str(search or "").strip()
        if not search_value and isinstance(filters, dict):
            search_value = str(filters.get("_search") or "").strip()
        if search_value:
            term = f"%{search_value.lower()}%"
            searchable = [func.lower(Record.data[col_id].astext).like(term) for col_id in columns]
            if searchable:
                stmt = stmt.where(or_(*searchable))

        count_stmt = select(func.count()).select_from(stmt.subquery())
        total = int((await self.session.execute(count_stmt)).scalar() or 0)

        order_clauses = self._build_sort_clauses(sorts=sorts, columns=columns)
        if order_clauses:
            stmt = stmt.order_by(*order_clauses)
        else:
            stmt = stmt.order_by(Record.position.asc(), Record.created_at.desc())

        page_stmt = stmt.limit(limit).offset(offset)
        page = list((await self.session.execute(page_stmt)).scalars().all())
        await self.relation_engine.enrich_records_for_read(table=table, records=page)
        await self.formula_engine.enrich_records_for_read(table=table, records=page)
        return {"records": page, "total": total}

    async def export_csv(
        self,
        *,
        table_id: uuid.UUID,
        org_id: uuid.UUID,
    ) -> tuple[bytes, str, str]:
        """Экспорт таблицы в CSV.

        Возвращает: (payload, media_type, filename)
        """
        table = await self.t_repo.get_by_id(table_id, with_columns=True)
        if not table or table.org_id != org_id:
            raise LookupError("NOT_FOUND")

        records = await self.r_repo.list_by_table(table_id, limit=5000, offset=0)
        columns = sorted(table.columns, key=lambda c: c.position)
        await self.relation_engine.enrich_records_for_read(table=table, records=records)
        await self.formula_engine.enrich_records_for_read(table=table, records=records)
        max_columns = int(max(1, settings.TABLE_EXPORT_MAX_COLUMNS))
        if len(columns) > max_columns:
            raise ValueError("EXPORT_TOO_MANY_COLUMNS")
        max_rows = int(max(1, settings.TABLE_EXPORT_MAX_ROWS))
        if len(records) > max_rows:
            raise ValueError("EXPORT_TOO_MANY_ROWS")

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([c.name for c in columns])
        for rec in records[:max_rows]:
            writer.writerow([str(rec.data.get(str(c.id), "")) for c in columns])

        payload = output.getvalue().encode("utf-8-sig")
        with suppress(Exception):
            EXPORTS_TOTAL.labels(format="csv").inc()

        return payload, "text/csv; charset=utf-8", f"{table.name}.csv"

    async def export_xlsx(
        self,
        *,
        table_id: uuid.UUID,
        org_id: uuid.UUID,
    ) -> tuple[bytes, str, str]:
        """Экспорт таблицы в XLSX.

        Возвращает: (payload, media_type, filename)
        """
        table = await self.t_repo.get_by_id(table_id, with_columns=True)
        if not table or table.org_id != org_id:
            raise LookupError("NOT_FOUND")

        records = await self.r_repo.list_by_table(table_id, limit=5000, offset=0)
        columns = sorted(table.columns, key=lambda c: c.position)
        await self.relation_engine.enrich_records_for_read(table=table, records=records)
        await self.formula_engine.enrich_records_for_read(table=table, records=records)
        max_columns = int(max(1, settings.TABLE_EXPORT_MAX_COLUMNS))
        if len(columns) > max_columns:
            raise ValueError("EXPORT_TOO_MANY_COLUMNS")
        max_rows = int(max(1, settings.TABLE_EXPORT_MAX_ROWS))
        if len(records) > max_rows:
            raise ValueError("EXPORT_TOO_MANY_ROWS")

        wb = Workbook()
        ws = wb.active
        ws.title = "Table"
        ws.append([c.name for c in columns])
        for rec in records[:max_rows]:
            ws.append([str(rec.data.get(str(c.id), "")) for c in columns])

        output = BytesIO()
        wb.save(output)
        payload = output.getvalue()

        with suppress(Exception):
            EXPORTS_TOTAL.labels(format="xlsx").inc()

        return payload, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", f"{table.name}.xlsx"

    async def import_csv(
        self,
        *,
        table_id: uuid.UUID,
        org_id: uuid.UUID,
        user_id: uuid.UUID,
        raw_bytes: bytes,
        mode: str = "append",
    ) -> dict:
        """Legacy single-step import. Kept for backward-compat."""
        result = await self.commit_csv_import(
            table_id=table_id,
            org_id=org_id,
            user_id=user_id,
            raw_bytes=raw_bytes,
            mode=mode,
            mapping_json=None,
            strict=True,
        )
        return {
            "created": int(result.get("records_created", 0)),
            "deleted_before": int(result.get("deleted_before", 0)),
        }

    async def preview_csv_import(
        self,
        *,
        table_id: uuid.UUID,
        org_id: uuid.UUID,
        raw_bytes: bytes,
        mode: str = "append",
        mapping_json: str | None = None,
    ) -> dict:
        table = await self.t_repo.get_by_id(table_id, with_columns=True)
        if not table or table.org_id != org_id:
            raise LookupError("NOT_FOUND")
        if len(raw_bytes) > int(max(1, settings.TABLE_IMPORT_MAX_BYTES)):
            raise ValueError("CSV_TOO_LARGE")

        mode_normalized = (mode or "append").strip().lower()
        if mode_normalized not in {"append", "replace"}:
            raise ValueError("BAD_IMPORT_MODE")

        parsed = self._parse_csv_rows(
            table=table,
            raw_bytes=raw_bytes,
            mapping_json=mapping_json,
        )
        return {
            "mode": mode_normalized,
            "header": parsed["header"],
            "matched_columns": parsed["matches"],
            "total_rows": int(parsed["total_rows"]),
            "valid_rows": len(parsed["rows"]),
            "invalid_rows": int(parsed["invalid_rows"]),
            "sample_rows": parsed["rows"][:20],
            "errors": parsed["errors"],
        }

    async def commit_csv_import(
        self,
        *,
        table_id: uuid.UUID,
        org_id: uuid.UUID,
        user_id: uuid.UUID,
        raw_bytes: bytes,
        mode: str = "append",
        mapping_json: str | None = None,
        strict: bool = False,
    ) -> dict:
        table = await self.t_repo.get_by_id(table_id, with_columns=True)
        if not table or table.org_id != org_id:
            raise LookupError("NOT_FOUND")
        if len(raw_bytes) > int(max(1, settings.TABLE_IMPORT_MAX_BYTES)):
            raise ValueError("CSV_TOO_LARGE")

        mode_normalized = (mode or "append").strip().lower()
        if mode_normalized not in {"append", "replace"}:
            raise ValueError("BAD_IMPORT_MODE")

        parsed = self._parse_csv_rows(
            table=table,
            raw_bytes=raw_bytes,
            mapping_json=mapping_json,
        )
        if strict and parsed["invalid_rows"] > 0:
            raise ValueError("CSV_VALIDATION_FAILED")

        valid_rows: list[dict] = parsed["rows"]
        to_create_count = len(valid_rows)
        deleted_before = 0
        if mode_normalized == "replace":
            existing_records = await self.r_repo.list_by_table(table_id=table_id, limit=100000, offset=0)
            for rec in existing_records:
                await self._append_audit_event(
                    org_id=org_id,
                    actor_id=user_id,
                    action=AuditAction.DELETE,
                    table_id=table_id,
                    record_id=rec.id,
                    before_data=dict(rec.data or {}),
                    after_data=None,
                    source="records.import_csv_replace_delete",
                    changed_columns=[],
                    extra_meta={"import_mode": mode_normalized},
                )
        if mode_normalized == "replace" and (to_create_count > 0 or parsed["total_rows"] == 0):
            deleted_before = await self.r_repo.delete_by_table(table_id=table_id)
        if to_create_count <= 0:
            return {
                "mode": mode_normalized,
                "records_created": 0,
                "records_skipped": int(parsed["invalid_rows"]),
                "deleted_before": deleted_before,
                "total_rows": int(parsed["total_rows"]),
                "errors": parsed["errors"],
            }

        await self._ensure_record_capacity(org_id=org_id, incoming_count=to_create_count)
        max_pos = await self.r_repo.get_max_position(table_id)
        rows_to_persist = [
            Record(
                table_id=table_id,
                org_id=org_id,
                created_by=user_id,
                data=row["data"],
                position=max_pos + idx + 1,
            )
            for idx, row in enumerate(valid_rows)
        ]
        await self.r_repo.bulk_create(rows_to_persist)
        for rec in rows_to_persist:
            await self._append_audit_event(
                org_id=org_id,
                actor_id=user_id,
                action=AuditAction.CREATE,
                table_id=table_id,
                record_id=rec.id,
                before_data=None,
                after_data=dict(rec.data or {}),
                source="records.import_csv_commit",
                changed_columns=sorted([str(k) for k in (rec.data or {})]),
                extra_meta={"import_mode": mode_normalized},
            )

        with suppress(Exception):
            IMPORTS_TOTAL.labels(format="csv").inc()
        return {
            "mode": mode_normalized,
            "records_created": int(to_create_count),
            "records_skipped": int(parsed["invalid_rows"]),
            "deleted_before": int(deleted_before),
            "total_rows": int(parsed["total_rows"]),
            "errors": parsed["errors"],
        }

    def _parse_csv_rows(
        self,
        *,
        table,
        raw_bytes: bytes,
        mapping_json: str | None,
    ) -> dict:
        started = time.monotonic()
        max_processing_s = float(max(0.1, settings.TABLE_IMPORT_MAX_PROCESSING_S))
        max_rows = int(max(1, settings.TABLE_IMPORT_MAX_ROWS))
        max_cols = int(max(1, settings.TABLE_IMPORT_MAX_COLUMNS))
        max_cell_chars = int(max(1, settings.TABLE_IMPORT_MAX_CELL_CHARS))

        text = raw_bytes.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        try:
            header = next(reader)
        except StopIteration as exc:
            raise ValueError("EMPTY_CSV") from exc
        if len(header) > max_cols:
            raise ValueError("TOO_MANY_COLUMNS")

        mapping = self._parse_mapping_json(mapping_json)
        mapped_columns = self._build_csv_mapping(table=table, header=header, mapping=mapping)
        if not mapped_columns:
            raise ValueError("NO_MATCHING_COLUMNS")

        rows: list[dict] = []
        errors: list[dict] = []
        invalid_rows = 0
        total_rows = 0
        for row_idx, row in enumerate(reader, start=2):
            if (time.monotonic() - started) > max_processing_s:
                raise ValueError("IMPORT_TIMEOUT")
            total_rows += 1
            if total_rows > max_rows:
                raise ValueError("TOO_MANY_ROWS")

            row_errors: list[dict] = []
            out_data: dict[str, object] = {}
            for item in mapped_columns:
                idx = int(item["csv_index"])
                col = item["column"]
                raw = str(row[idx]).strip() if idx < len(row) and row[idx] is not None else ""
                if raw and len(raw) > max_cell_chars:
                    row_errors.append(
                        {
                            "row_number": row_idx,
                            "column": col.name,
                            "code": "CELL_TOO_LARGE",
                            "message": "Слишком длинное значение",
                            "raw_value": raw[:80],
                        }
                    )
                    continue
                if not raw:
                    if col.is_required:
                        row_errors.append(
                            {
                                "row_number": row_idx,
                                "column": col.name,
                                "code": "REQUIRED",
                                "message": "Обязательное поле пустое",
                                "raw_value": None,
                            }
                        )
                    continue

                normalized, err = self._normalize_cell_value(field_type=str(col.field_type), raw=raw)
                if err:
                    row_errors.append(
                        {
                            "row_number": row_idx,
                            "column": col.name,
                            "code": err,
                            "message": f"Некорректное значение для типа {col.field_type}",
                            "raw_value": raw,
                        }
                    )
                    continue
                out_data[str(col.id)] = normalized

            if row_errors:
                invalid_rows += 1
                errors.extend(row_errors[:10])
                continue
            if out_data:
                rows.append({"row_number": row_idx, "data": out_data})

        matches_out = [
            {
                "csv_column": str(item["csv_name"]),
                "table_column_id": str(item["column"].id),
                "table_column_name": str(item["column"].name),
            }
            for item in mapped_columns
        ]
        return {
            "header": [str(x) for x in header],
            "matches": matches_out,
            "rows": rows,
            "errors": errors[:200],
            "invalid_rows": invalid_rows,
            "total_rows": total_rows,
        }

    @staticmethod
    def _parse_mapping_json(mapping_json: str | None) -> dict[str, str]:
        if not mapping_json:
            return {}
        try:
            raw = json.loads(mapping_json)
        except Exception as exc:
            raise ValueError("BAD_MAPPING_JSON") from exc
        if not isinstance(raw, dict):
            raise ValueError("BAD_MAPPING_JSON")
        out: dict[str, str] = {}
        for key, value in raw.items():
            k = str(key).strip()
            v = str(value).strip()
            if k and v:
                out[k] = v
        return out

    @staticmethod
    def _build_csv_mapping(*, table, header: list[str], mapping: dict[str, str]) -> list[dict]:
        by_name = {str(c.name).strip().lower(): c for c in table.columns}
        by_id = {str(c.id): c for c in table.columns}
        out: list[dict] = []
        for idx, source_name in enumerate(header):
            source_name_clean = str(source_name).strip()
            mapped_to = mapping.get(source_name_clean) or mapping.get(str(idx))
            target_col = None
            if mapped_to:
                target_col = by_id.get(mapped_to) or by_name.get(mapped_to.lower())
            if target_col is None:
                target_col = by_name.get(source_name_clean.lower())
            if target_col is None:
                continue
            out.append({"csv_index": idx, "csv_name": source_name_clean, "column": target_col})
        return out

    @staticmethod
    def _normalize_cell_value(*, field_type: str, raw: str) -> tuple[object, str | None]:
        ft = str(field_type).strip().lower()
        if ft == "number":
            try:
                return float(raw), None
            except Exception:
                return raw, "INVALID_NUMBER"
        if ft == "boolean":
            value = raw.strip().lower()
            if value in {"true", "1", "yes", "y", "да"}:
                return True, None
            if value in {"false", "0", "no", "n", "нет"}:
                return False, None
            return raw, "INVALID_BOOLEAN"
        if ft == "date":
            try:
                d = date.fromisoformat(raw.strip())
                return d.isoformat(), None
            except Exception:
                return raw, "INVALID_DATE"
        if ft == "datetime":
            try:
                dt = datetime.fromisoformat(raw.strip())
                return dt.isoformat(), None
            except Exception:
                return raw, "INVALID_DATETIME"
        if ft == "multi_select":
            values = [part.strip() for part in raw.split(",") if part.strip()]
            return values, None
        return raw, None

    async def _ensure_record_capacity(self, *, org_id: uuid.UUID, incoming_count: int) -> None:
        if incoming_count <= 0:
            return
        plan = await self.plan_limits_repo.resolve_effective_plan(org_id=org_id)
        limit = int(getattr(plan, "max_records", 0) or 0)
        if limit <= 0:
            return
        current = await self.plan_limits_repo.count_records_by_org(org_id)
        if current + incoming_count > limit:
            raise ValueError("RECORD_LIMIT_REACHED")

    @staticmethod
    def _normalize_filters(filters: dict | list | None) -> list[dict]:
        if not filters:
            return []
        if isinstance(filters, list):
            out: list[dict] = []
            for item in filters:
                if hasattr(item, "model_dump"):
                    item = item.model_dump()
                if isinstance(item, dict) and item.get("col_id"):
                    out.append(item)
            return out
        if isinstance(filters, dict):
            out = []
            for col_id, cond in filters.items():
                if col_id == "_search":
                    continue
                if isinstance(cond, dict):
                    out.append({"col_id": col_id, "op": cond.get("op", "eq"), "value": cond.get("value")})
                else:
                    out.append({"col_id": col_id, "op": "eq", "value": cond})
            return out
        return []

    @staticmethod
    def _build_filter_condition(*, item: dict, columns: dict) -> object | None:
        col_id = str(item.get("col_id") or "").strip()
        if not col_id or col_id not in columns:
            return None
        op = str(item.get("op") or "contains").strip().lower()
        value = item.get("value")
        field_type = str(columns[col_id].field_type)
        cell_text = Record.data[col_id].astext

        if op == "is_empty":
            return or_(cell_text.is_(None), cell_text == "")

        if op == "contains":
            needle = str(value or "").strip().lower()
            if not needle:
                return None
            return func.lower(cell_text).like(f"%{needle}%")

        if op == "in":
            values = value if isinstance(value, list) else [value]
            normalized = [str(v) for v in values if v is not None and str(v) != ""]
            if not normalized:
                return None
            return cell_text.in_(normalized)

        if op == "between":
            left = right = None
            if isinstance(value, dict):
                left = value.get("from")
                right = value.get("to")
            elif isinstance(value, list) and len(value) >= 2:
                left, right = value[0], value[1]
            if left is None or right is None:
                return None
            if field_type == "number":
                try:
                    left_num = float(left)
                    right_num = float(right)
                except (TypeError, ValueError):
                    return None
                return and_(cast(cell_text, Float) >= left_num, cast(cell_text, Float) <= right_num)
            if field_type == "date":
                return and_(
                    cast(cell_text, Date) >= cast(str(left), Date),
                    cast(cell_text, Date) <= cast(str(right), Date),
                )
            if field_type == "datetime":
                return and_(
                    cast(cell_text, DateTime) >= cast(str(left), DateTime),
                    cast(cell_text, DateTime) <= cast(str(right), DateTime),
                )
            return and_(cell_text >= str(left), cell_text <= str(right))

        if value is None:
            return None

        if field_type == "number":
            try:
                rhs = float(value)
            except (TypeError, ValueError):
                return None
            lhs = cast(cell_text, Float)
        elif field_type == "date":
            lhs = cast(cell_text, Date)
            rhs = cast(str(value), Date)
        elif field_type == "datetime":
            lhs = cast(cell_text, DateTime)
            rhs = cast(str(value), DateTime)
        else:
            lhs = cell_text
            rhs = str(value)

        if op == "eq":
            return lhs == rhs
        if op == "neq":
            return lhs != rhs
        if op == "gt":
            return lhs > rhs
        if op == "lt":
            return lhs < rhs
        return None

    @staticmethod
    def _build_sort_clauses(*, sorts: list | None, columns: dict) -> list[object]:
        if not sorts:
            return []
        clauses: list[object] = []
        for s in sorts:
            if hasattr(s, "model_dump"):
                s = s.model_dump()
            if not isinstance(s, dict):
                continue
            col_id = str(s.get("col_id") or "").strip()
            if not col_id or col_id not in columns:
                continue
            desc = str(s.get("dir") or "asc").lower() == "desc"
            field_type = str(columns[col_id].field_type)
            cell_text = Record.data[col_id].astext
            if field_type == "number":
                expr = cast(cell_text, Float)
            elif field_type == "date":
                expr = cast(cell_text, Date)
            elif field_type == "datetime":
                expr = cast(cell_text, DateTime)
            else:
                expr = cell_text
            clauses.append(expr.desc() if desc else expr.asc())
        return clauses

    async def _append_audit_event(
        self,
        *,
        org_id: uuid.UUID,
        actor_id: uuid.UUID | None,
        action: AuditAction,
        table_id: uuid.UUID,
        record_id: uuid.UUID,
        before_data: dict | None,
        after_data: dict | None,
        source: str,
        changed_columns: list[str],
        extra_meta: dict | None = None,
    ) -> None:
        meta = {
            "table_id": str(table_id),
            "record_id": str(record_id),
            "source": source,
            "changed_columns": changed_columns,
            "before_data": before_data,
            "after_data": after_data,
        }
        if extra_meta:
            meta.update(extra_meta)
        self.session.add(
            AuditLog(
                org_id=org_id,
                actor_id=actor_id,
                action=action,
                entity_type="table_record",
                entity_id=str(record_id),
                meta=meta,
            )
        )
        await self.session.flush()
